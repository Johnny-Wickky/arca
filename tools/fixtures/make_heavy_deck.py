# SPDX-License-Identifier: AGPL-3.0-only
"""Generates a heavy picture deck (about 12 MB) for the "Pictures in the text" checks.

Needs python-pptx and Pillow:

    python3 tools/fixtures/make_heavy_deck.py out/heavy.pptx [slides] [pictures-per-slide]

Every slide has a picture background (`p:bg/p:bgPr/a:blipFill`), a few `p:pic` photos and one
shape filled with a picture (`a:blipFill` in `p:spPr`); the master has a full-bleed `p:pic` on
the back layer. The pictures are noise, so they do not compress and the file stays large.
All content is synthetic. Nothing here is, or contains, malware.
"""
import io
import random
import sys

from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def noise(rng, w, h, tint):
    """A noisy JPEG (poorly compressible) with a dominant colour, so pictures can be told apart."""
    raw = bytes(rng.randrange(256) for _ in range(w * h * 3))
    im = Image.frombytes("RGB", (w, h), raw)
    solid = Image.new("RGB", (w, h), tint)
    im = Image.blend(im, solid, 0.55)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    return buf.getvalue()


def picture_background(slide, rid):
    """Gives the slide `<p:bg><p:bgPr><a:blipFill>` with the picture `rid`."""
    c_sld = slide._element.find(f"{{{NS_P}}}cSld")
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:bgPr>'
        f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
        f'<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'
    )
    c_sld.insert(0, bg)


def picture_filled_shape(slide, rid, x, y, w, h):
    """A rectangle whose fill is the picture `rid` (`a:blipFill` inside `p:spPr`)."""
    shape = slide.shapes.add_shape(1, x, y, w, h)
    sp_pr = shape._element.spPr
    for child in list(sp_pr):
        if child.tag in (f"{{{NS_A}}}solidFill", f"{{{NS_A}}}noFill"):
            sp_pr.remove(child)
    fill = etree.fromstring(
        f'<a:blipFill xmlns:a="{NS_A}" xmlns:r="{NS_R}"><a:blip r:embed="{rid}"/>'
        f'<a:stretch><a:fillRect/></a:stretch></a:blipFill>'
    )
    geom = sp_pr.find(f"{{{NS_A}}}prstGeom")
    geom.addnext(fill)


def build(path, slides=24, pics=3):
    rng = random.Random(7)
    prs = Presentation()
    master = prs.slide_master
    # The master carries a full-bleed picture on the back layer.
    _, mrid = master.part.get_or_add_image_part(io.BytesIO(noise(rng, 640, 480, (30, 60, 120))))
    tree = master.shapes._spTree
    pic = etree.fromstring(
        f'<p:pic xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:nvPicPr>'
        f'<p:cNvPr id="900" name="Master backdrop" descr="backdrop"/><p:cNvPicPr/><p:nvPr userDrawn="1"/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{mrid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="9144000" cy="6858000"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
    )
    tree.insert(2, pic)
    for i in range(slides):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Slide {i + 1}"
        tint = (rng.randrange(256), rng.randrange(256), rng.randrange(256))
        _, bg_rid = s.part.get_or_add_image_part(io.BytesIO(noise(rng, 800, 600, tint)))
        picture_background(s, bg_rid)
        for j in range(pics):
            data = noise(rng, 480, 360, tint)
            s.shapes.add_picture(
                io.BytesIO(data), Inches(0.4 + 3.0 * (j % 3)), Inches(1.8 + 2.4 * (j // 3)), Inches(2.8), Inches(2.1)
            )
        _, fill_rid = s.part.get_or_add_image_part(io.BytesIO(noise(rng, 320, 240, tint)))
        picture_filled_shape(s, fill_rid, Inches(0.4), Inches(6.3), Inches(2.0), Inches(0.9))
    prs.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "heavy.pptx"
    n = int(sys.argv[2]) if len(sys.argv) > 2 else 24
    k = int(sys.argv[3]) if len(sys.argv) > 3 else 3
    build(out, n, k)
