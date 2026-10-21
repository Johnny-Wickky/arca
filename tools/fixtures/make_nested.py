#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-only
"""Writes a .docx with a table inside a table and numbered lists of two levels; convert it with

    soffice --headless --convert-to doc:"MS Word 97" nested.docx

to get tests/fixtures/nested.doc (a Word 97 file whose nested cells and rows are marked the way
Word marks them: itap, sprmPFInnerTableCell, sprmPFInnerTtp).
"""
import sys

import docx

d = docx.Document()
d.add_heading("Вложенные таблицы и списки", 1)
t = d.add_table(rows=2, cols=3)
t.style = "Table Grid"
for r in range(2):
    for c in range(3):
        t.cell(r, c).text = f"O{r}{c}"
inner = t.cell(1, 1).add_table(rows=2, cols=2)
inner.style = "Table Grid"
for r in range(2):
    for c in range(2):
        inner.cell(r, c).text = f"I{r}{c}"
t.cell(1, 1).add_paragraph("после вложенной")
d.add_paragraph("A", style="List Number")
d.add_paragraph("B", style="List Number")
d.add_paragraph("C", style="List Number 2")
d.add_paragraph("D", style="List Number 2")
d.add_paragraph("E", style="List Number")
d.save(sys.argv[1] if len(sys.argv) > 1 else "nested.docx")
