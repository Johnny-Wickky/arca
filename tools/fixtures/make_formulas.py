#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-only
"""Writes tests/fixtures/formulas.xlsx: a workbook of the kind a library produces — formulas and
no stored results (openpyxl writes none and asks the application to recalculate on open).

    pip install openpyxl
    python3 tools/fixtures/make_formulas.py tests/fixtures/formulas.xlsx

Nothing in it is calculated here: the point of the file is that the viewer has to.
"""
import sys

import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Продажи"
ws["A1"], ws["B1"] = "Товар", "Сумма"
for i, (name, amount) in enumerate([("Яблоки", 40), ("Груши", 25.5), ("Сливы", 34.5)], start=2):
    ws.cell(i, 1, name)
    ws.cell(i, 2, amount)
ws["A5"] = "Итого"
ws["B5"] = "=SUM(B2:B4)"
ws["A6"] = "Подпись"
ws["B6"] = '="Всего: "&TEXT(B5,"0.00")&" руб."'
ws["A7"] = "Больше всех"
ws["B7"] = "=INDEX(A2:A4,MATCH(MAX(B2:B4),B2:B4,0))"
ws["A8"] = "Доля яблок"
ws["B8"] = "=B2/B5"
ws["B8"].number_format = "0%"
ws["A9"] = "Сегодня"
ws["B9"] = "=TODAY()"  # needs the clock: stays a formula

other = wb.create_sheet("Сводка")
other["A1"] = "Из первого листа"
other["B1"] = "=Продажи!B5*2"
wb.save(sys.argv[1] if len(sys.argv) > 1 else "formulas.xlsx")
